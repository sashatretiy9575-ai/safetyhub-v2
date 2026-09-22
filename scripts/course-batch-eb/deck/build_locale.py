# -*- coding: utf-8 -*-
"""Builds the electrical-safety deck in another language with the owner's own
generator: the Russian slide data is replaced by its translation and the few
strings the generator writes itself are taken from ui-<locale>.json. Nothing
about the layout changes.

The translations live in `content/course-batch-2026-09-eb/deck-source`; the
photographs and the logo are the owner's own and stay out of the repository, so
the script is run from a working directory that holds `assets/`.

Usage: SOURCE=<dir with slides-*.json> python build_locale.py kk|en|zh
       -> deck-<locale>.pptx
"""
import json
import re
import sys
import types

import os

LOCALE = sys.argv[1]
SOURCE = os.environ.get("SOURCE", ".")
slides = json.load(open(os.path.join(SOURCE, f"slides-{LOCALE}.json"), encoding="utf-8"))
ui = json.load(open(os.path.join(SOURCE, f"ui-{LOCALE}.json"), encoding="utf-8"))

source = open("build_safetyhub_deck.py", encoding="utf-8").read()
replacements = [
    ('f"Норма: {norm}"', 'f"{UI[\'norm_prefix\']} {norm}"'),
    ('"Электробезопасность (ЭБ) и работы на высоте • Практический курс SafetyHUB"', "UI['cover_footer']"),
    ('"Электробезопасность (ЭБ) и работы на высоте • SafetyHUB"', "UI['footer']"),
    ('p_oct.text = "СТОП"', "p_oct.text = UI['stop']"),
    ('if ":" in bullet_text:', 'if _label_split(bullet_text):'),
    ('parts = bullet_text.split(":", 1)', 'parts = _label_split(bullet_text)'),
    ('prefix = parts[0] + ":"', 'prefix = parts[0] + parts[2]'),
]
for old, new in replacements:
    if old not in source:
        raise SystemExit("generator changed, not found: " + old)
    source = source.replace(old, new)
if re.search(r'"[^"\n]*[А-Яа-яЁё][^"\n]*"', source.split('def generate_presentation')[0].split('"""', 2)[2]):
    for line in source.splitlines():
        if re.search(r'"[^"\n]*[А-Яа-яЁё][^"\n]*"', line) and 'data.get(' not in line and 'def generate_presentation' not in line and not line.strip().startswith('#'):
            print("warning: Russian literal left:", line.strip())

if LOCALE == "zh":
    source = source.replace('FONT_FAMILY = "Arial"', 'FONT_FAMILY = "Microsoft YaHei"')
    # A latin typeface alone does not reach Chinese characters: PowerPoint
    # draws them with the East Asian face, which must be named too.
    from lxml import etree
    from pptx.oxml.ns import qn
    from pptx.text.text import Font

    setter = Font.name.fset

    def name_with_east_asian(self, value):
        setter(self, value)
        rPr = self._rPr
        # PowerPoint picks the line-breaking rules from the run's language: a
        # run left in English lets a Chinese comma open a line.
        rPr.set("lang", "zh-CN")
        rPr.set("altLang", "en-US")
        latin = rPr.find(qn("a:latin"))
        previous = latin
        for tag in ("a:ea", "a:cs"):
            element = rPr.find(qn(tag))
            if element is None:
                element = etree.SubElement(rPr, qn(tag))
                previous.addnext(element)
            element.set("typeface", value)
            previous = element

    Font.name = property(Font.name.fget, name_with_east_asian)

    # Chinese line breaking: a closing bracket, a full stop or an exclamation
    # mark must not open a line, and an opening bracket must not end one.
    from pptx.text.text import _Paragraph

    paragraph_text = _Paragraph.text.fset

    def kinsoku(paragraph):
        pPr = paragraph._p.get_or_add_pPr()
        pPr.set("eaLnBrk", "1")
        pPr.set("hangingPunct", "1")
        pPr.set("latinLnBrk", "0")

    def text_with_kinsoku(self, value):
        paragraph_text(self, value)
        kinsoku(self)

    _Paragraph.text = property(_Paragraph.text.fget, text_with_kinsoku)
    add_run = _Paragraph.add_run

    def run_with_kinsoku(self):
        kinsoku(self)
        return add_run(self)

    _Paragraph.add_run = run_with_kinsoku

def _balanced(text, lines):
    """Hands a hand-broken line to PowerPoint as running text, one sentence a line.

    The breaks in the Russian deck were placed by eye for Russian words at
    Russian widths. Re-using those positions in another language strands a word
    or a dash on a line of its own, and re-breaking by character count guesses
    at widths PowerPoint alone knows. So each sentence is given as one line and
    wrapped by the program that measures it, and a new statement still starts a
    line of its own, as it does on the Russian page.
    """
    running = " ".join(text.replace("\n", " ").split())
    sentences = [part.strip() for part in re.split(r"(?<=[.!?]) +", running) if part.strip()]
    return "\n".join(sentences) if len(sentences) > 1 else running


def _rewrap(value, source):
    """Every hand-broken string of the translation, re-broken for its own words."""
    if isinstance(value, str) and isinstance(source, str) and "\n" in source:
        return _balanced(value, source.count("\n") + 1)
    if isinstance(value, list) and isinstance(source, list) and len(value) == len(source):
        return [_rewrap(item, source[index]) for index, item in enumerate(value)]
    if isinstance(value, dict) and isinstance(source, dict):
        return {key: _rewrap(item, source.get(key)) for key, item in value.items()}
    return value


# Chinese is broken by the translator at its own measure: it has no spaces to
# re-break at, and a line may not part a number from its unit.
if LOCALE not in ("ru", "zh"):
    russian = json.load(open(os.path.join(SOURCE, "slides-ru.json"), encoding="utf-8"))
    slides = [_rewrap(slide, russian[index]) for index, slide in enumerate(slides)]

deck_content = types.ModuleType("deck_content")
deck_content.SLIDES_DATA = slides
sys.modules["deck_content"] = deck_content
module = types.ModuleType("build_localized")
module.UI = ui


def _label_split(text):
    match = re.search(r"[:：]", text)
    if not match:
        return None
    return [text[: match.start()], text[match.end():], match.group(0)]


module._label_split = _label_split
exec(compile(source, "build_safetyhub_deck.py", "exec"), module.__dict__)
module.generate_presentation(f"deck-{LOCALE}.pptx")
